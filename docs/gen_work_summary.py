#!/usr/bin/env python3
"""生成《巡检机器人自主导航系统 工作总结》Word 文档"""
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

OUT = "/home/ubuntu/my_ros2_ws/docs/工作总结_巡检机器人自主导航系统.docx"

doc = Document()

# ---------- 页面与默认样式 ----------
sec = doc.sections[0]
sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
sec.left_margin = sec.right_margin = Cm(2.5)
sec.top_margin = sec.bottom_margin = Cm(2.5)

normal = doc.styles["Normal"]
normal.font.name = "Times New Roman"
normal.font.size = Pt(11)
normal._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")

for name, size in (("Heading 1", 14), ("Heading 2", 12), ("Heading 3", 11)):
    st = doc.styles[name]
    st.font.name = "Times New Roman"
    st.font.size = Pt(size)
    st.font.bold = True
    st.font.color.rgb = RGBColor(0, 0, 0)
    st._element.rPr.rFonts.set(qn("w:eastAsia"), "黑体")
    st.paragraph_format.space_before = Pt(10)
    st.paragraph_format.space_after = Pt(6)


def set_run(run, ea="宋体", size=11, bold=False):
    run.font.name = "Times New Roman"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), ea)
    run.font.size = Pt(size)
    run.font.bold = bold


def para(text, size=11, bold=False, align=None, space_after=6, ea="宋体"):
    p = doc.add_paragraph()
    r = p.add_run(text)
    set_run(r, ea=ea, size=size, bold=bold)
    p.paragraph_format.line_spacing = 1.3
    p.paragraph_format.space_after = Pt(space_after)
    if align is not None:
        p.alignment = align
    return p


def shade(cell, color="D9E2F3"):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:fill"), color)
    tcPr.append(shd)


def add_table(headers, rows, widths=None, font_size=10):
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = "Table Grid"
    t.autofit = False
    for i, h in enumerate(headers):
        cell = t.rows[0].cells[i]
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(h)
        set_run(r, ea="黑体", size=font_size, bold=True)
        shade(cell)
        if widths:
            cell.width = Cm(widths[i])
    for row in rows:
        cells = t.add_row().cells
        for i, v in enumerate(row):
            p = cells[i].paragraphs[0]
            r = p.add_run(str(v))
            set_run(r, size=font_size)
            if widths:
                cells[i].width = Cm(widths[i])
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return t


# ---------- 标题 ----------
para("巡检机器人自主导航系统 工作总结", size=16, bold=True, ea="黑体",
     align=WD_ALIGN_PARAGRAPH.CENTER, space_after=2)
para("2026 年 8 月", size=11, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=12)

# ---------- 一、项目概述 ----------
doc.add_heading("一、项目概述", level=1)
para("本项目为基于树莓派 5 的自主巡检小车(Hiwonder 平台),软件栈运行于 ROS2 Humble 容器内。"
     "硬件组成:MS200 2D 激光雷达(10Hz, 360°)、QMI8658 IMU、STM32 差分底盘;"
     "软件组成:robot_driver 底盘串口驱动、EKF/ESKF 里程计融合、Nav2 导航栈"
     "(AMCL 定位 + NavFn 全局规划 + DWB 局部控制 + velocity_smoother)、slam_toolbox 激光建图。")
para("项目实现了「记录航点 → 依次巡检 → 自动回原点」的完整自主巡检闭环。"
     "2026-08-25 实测:两航点完整巡检 2/2 通过,最终停点误差 0.12m / 3.3°。")

# ---------- 二、承担工作 ----------
doc.add_heading("二、本人承担的工作内容", level=1)
items = [
    "ESKF 里程计融合:从零完成 ESKF 学习项目(stage0~6)并落地为 ROS2 节点,与 robot_localization EKF 并列可选的定位方案;",
    "Nav2 自主巡检工作流:封装一键启动、航点记录/巡检脚本、全局定位与收敛判定的完整流程;",
    "SLAM 地图重建:定位玻璃墙干扰根因,三次建图最终得到零空洞的正式地图;",
    "到点判定疑难问题定位与修复:定位冻结、双重锁死、旋转死区三个根因全部解决并实测验证;",
    "底盘直行纠偏与参数整定:键盘遥控直行自动纠偏、速度与死区相关参数实测标定;",
    "开发环境与工程化:容器化部署、脚本分层、README 重写、git 提交整理与交接文档。",
]
for i, it in enumerate(items, 1):
    para(f"{i}. {it}")

# ---------- 三、主要工作与成果 ----------
doc.add_heading("三、主要工作与成果", level=1)

doc.add_heading("3.1 ESKF 里程计融合", level=2)
para("算法层面:完成 ESKF 学习项目 stage0~6,实现 15 维误差状态 ESKF"
     "(δp, δv, δθ, δb_g, δb_a),predict() 支持真实时间间隔,位置观测更新,共享核心代码位于 "
     "include/eskf.hpp。")
para("工程落地:开发 ROS2 节点 eskf_ros,订阅 /imu 与 /odom,发布 /odometry/filtered;"
     "以首条 /odom 的位置、朝向、速度完成初始化,此后 /odom 仅作位置观测;"
     "predict 在 IMU 回调、update 在里程计回调,按 header.stamp 时间序执行。")
para("仿真验证:纯积分 30m 距离漂移明显,引入 ESKF 后 ATE 0.33m / RPE 0.46m。"
     "实车接入真实话题(/imu/data、/odom),与 robot_localization EKF 并列可选,"
     "供定位方案对比评估。")

doc.add_heading("3.2 Nav2 自主巡检工作流", level=2)
para("一键启动:nav_bringup.launch.py 整合定位、规划、控制,地图路径参数化,"
     "默认使用正式图 maps/my_map.yaml。")
para("巡检主程序 waypoint_patrol.py:依次到达各航点、每点停驻(--dwell)、"
     "走完自动回原点(--home/--no-return-home)、单点超时取消(--waypoint-timeout)、"
     "循环巡检(--loops)。")
para("航点记录 record_waypoint.py:读取 /amcl_pose 记录航点;针对 AMCL 静止时不发布位姿的"
     "设计特性,加入 2 秒无位姿超时后转 tf 采样 map→base_link 的兜底机制,车停稳也能记录。")
para("全局定位方案:本版 AMCL(1.1.13)没有 /global_localization 服务,也不认零协方差 "
     "initialpose。等效做法是发布「大协方差 /initialpose」(均值取地图已探索区中心、"
     "协方差盖满全图),使粒子撒满全图、随车运动自动收敛,已封装为 global_localize.py;"
     "收敛判定 check_localization.py 以粒子云加权标准差 <0.2m 为锁定标志"
     "(粒子云为 best_effort QoS 的 nav2_msgs/ParticleCloud,订阅需用 qos_profile_sensor_data)。")

doc.add_heading("3.3 SLAM 地图重建", level=2)
para("问题定位:前两次建图失败根因是未走完整闭环,漂移不断累计导致地图被拉伸变乱;"
     "第三次沿外墙一口气走完一整圈、回到出发区域,回环闭合后地图自动拉齐。"
     "另外玻璃墙对雷达透射,全部贴白纸后重新建图。")
para("成果:新图 304×202px、分辨率 0.05m、探索区 12.8×7.9m、内部零空洞,"
     "升级为正式图并备份历史失败图;期间沉淀了建图操作与地图分析脚本(pgm 像素语义、"
     "连通域找洞、distance_transform 对齐评估等)。")

doc.add_heading("3.4 到点判定问题定位与修复(2026-08-25)", level=2)
para("现象:车到目标点约 0.15m 处停死,永远判不到达,触发 recovery 反复;到点后朝向对不齐。"
     "经排查定位到三个根因,全部修复并实测验证:")
add_table(
    ["序号", "根因", "修复", "参数变化"],
    [
        ["1", "AMCL 运动门控:目标点附近微调时定位冻结,不发布位姿",
         "调小更新门控", "update_min_d 0.25→0.05\nupdate_min_a 0.2→0.1"],
        ["2", "DWB 内部到点容差与 goal_checker 相等,车在 0.15m 处停平移只旋转,永远进不了判定圈(双重锁死)",
         "DWB 内部容差必须小于判定容差", "FollowPath.xy_goal_tolerance\n0.15→0.05"],
        ["3", "旋转死区:实测 wz=0.15 rad/s 不动(+0.06°),0.30 才正常转(+25°/1.5s),旧值 0.15 落在死区",
         "提高最小角速度下限", "FollowPath.min_speed_theta\n0.15→0.3"],
    ],
    widths=[1.2, 5.6, 4.2, 5.0],
)
para("最终到达容差(用户拍板):xy_goal_tolerance 0.15、yaw_goal_tolerance 0.12。")

doc.add_heading("3.5 直行纠偏与底盘参数整定", level=2)
para("键盘直行自动纠偏 straighten_teleop.py:检测纯直行命令后锁存参考直线(odom 系),"
     "10Hz 闭环输出「航向保持 PID + 横向纠偏」,转向/停止透传。"
     "关键结论:控制反馈必须用 odom 系——用 map 系(AMCL)反馈时,AMCL 移动中激光匹配跳变"
     "(map 位姿 2s 内可跳 20cm+),会误导纠偏导致「跑着跑着突然拐弯」。"
     "实测:3 段共 8m 直行,法向偏差 ≤1.9cm(目标 3cm),达标。")
para("底盘参数实测整定:电机死区约 0.1 m/s(低于该速度轮子不动),为 DWB 加 "
     "min_speed_xy 0.1 / min_speed_theta 0.15 保护(后按旋转死区实测调整到 0.3);"
     "最高速度 0.4→0.3 m/s(用户体验调优),到达容差 0.25→0.1→0.15(0.1 过严导致到点振荡,"
     "0.25 过松停点偏)。")

doc.add_heading("3.6 开发环境与工程化", level=2)
para("Docker 容器化部署(整仓挂载 + 串口 USB 直通),宿主机负责容器管理、容器内跑 ROS 栈,"
     "环境可复现;scripts/ 巡检工作流脚本与 tools/ 调试工具分层;README 重写为完整项目文档;"
     "git 三次提交整理全仓(首次全仓、到点三修复、文档整理);"
     "编写树莓派端交接文档(PI_HANDOFF.md)与学习札记(LEARNING_NOTES.md)。")

# ---------- 四、关键技术难点与解决方案 ----------
doc.add_heading("四、关键技术难点与解决方案", level=1)
add_table(
    ["难点", "解决方案"],
    [
        ["AMCL 车静止时不发布 /amcl_pose,航点记录等不到位姿",
         "record_waypoint.py 加 2s 超时,转 tf 采样 map→base_link(静止时持续发布,等效定位位姿)"],
        ["--initial-pose 发零协方差 initialpose,本版 AMCL 直接忽略",
         "改用 RViz 2D Pose Estimate 手动设置;或发布大协方差 /initialpose 撒满粒子做全局定位"],
        ["AMCL 粒子云「永远收不到」",
         "其为 best_effort QoS 的 nav2_msgs/ParticleCloud,须用 qos_profile_sensor_data 订阅;静止时不发布,需边动边看"],
        ["Nav2 取消操作挂起/下一个目标秒拒",
         "调底层服务 /navigate_to_pose/_action/cancel_goal(空 goal_info=取消全部);取消后 recovery 行为残留,加 3s 静置再发下一目标"],
        ["驱动进程被杀后 STM32 执行残留指令,车原地打转",
         "停车顺序:先取消 goal/发零速再停进程;救急:向 /dev/ttyACM0 连发 CMD,0,0,0"],
        ["键盘 teleop 未退出时 Nav2 下轮子冻结",
         "straighten_teleop 与 Nav2 抢 /cmd_vel,巡检前完全退出键盘,并用 ros2 topic info /cmd_vel -v 核对发布者"],
        ["搬车导致 AMCL 被「绑架」跳到错误位置,航点全废",
         "流程纪律:导航运行中绝不手搬车,挪车一律用键盘 teleop"],
        ["AMCL 定位不认 laser_frame_id 参数、排查定位不发位姿",
         "Humble 版 AMCL 用 scan 消息自带 frame_id;排查先看 /MS200/scan 与 tf 是否在"],
        ["容器内没有 docker CLI、pkill -f 会误杀自身 shell、pkill -x 对长进程名无效",
         "宿主机管理容器;pkill 用精确进程名/正则括号技巧([_] 防自匹配)/直接 kill PID"],
    ],
    widths=[6.5, 9.5],
)

# ---------- 五、验证结果 ----------
doc.add_heading("五、验证结果", level=1)
doc.add_heading("5.1 完整巡检(2026-08-25 实测)", level=2)
add_table(
    ["阶段", "目标点(地图系)", "用时", "过程特征"],
    [
        ["起点 → 航点1", "(5.696, 0.336, 32°)", "22.6s", "全程 5.71m、0.27m/s 匀速,距离单调收敛"],
        ["停驻", "—", "3s", "—"],
        ["航点1 → 航点2", "(6.789, 0.216, -4.8°)", "7.2s", "—"],
        ["停驻", "—", "3s", "—"],
        ["航点2 → 原点", "(0, 0, 0)", "45.0s", "中途掉头旋转 + AMCL 位姿跳变,自动恢复"],
    ],
    widths=[3.5, 5.0, 2.0, 5.5],
)
para("结果:两航点巡检 2/2 成功,最终停点误差 0.12m / 3.3°,判定逻辑稳定无振荡。")

doc.add_heading("5.2 直行纠偏", level=2)
para("键盘直行 3 段共 8m,法向偏差 ≤1.9cm(目标 3cm),达标。")
doc.add_heading("5.3 地图质量", level=2)
para("304×202px、0.05m 分辨率、探索区 12.8×7.9m、内部零空洞、回环闭合无拉伸。")
doc.add_heading("5.4 死区实测", level=2)
para("旋转:wz=0.15 rad/s 不动(+0.06°),0.30 正常转(+25°/1.5s);平移:低于 0.1 m/s 轮子不转。")

# ---------- 六、经验教训 ----------
doc.add_heading("六、经验教训", level=1)
items = [
    "先测硬件特性、再调软件参数:死区、容差这类参数不能靠猜,实测(旋转死区、最小速度)后一次调对;",
    "闭环思维:建图必须走完整闭环才能回环闭合;巡检本身也是一条闭环链路,环节之间互相依赖;",
    "控制反馈要用基准稳定的坐标系:odom 系(里程计)连续稳定,map 系(AMCL)存在匹配跳变,控制回路不能用后者;",
    "流程纪律比调参更值钱:不搬车、先取消目标再停进程、巡检前退出键盘,这些纪律避免了大量返工;",
    "定位层的「设计特性」不是 bug:AMCL 静止不发布、零协方差不认,都是设计行为,理解后针对性绕过即可。",
]
for i, it in enumerate(items, 1):
    para(f"{i}. {it}")

# ---------- 七、后续展望 ----------
doc.add_heading("七、后续展望", level=1)
items = [
    "视觉感知接入:引入双目/深度相机,补充低矮与悬空障碍检测,与 2D 雷达融合增强避障能力(当前避障仅依赖 2D 雷达);",
    "多航点、长时间巡检的规模测试,完善异常恢复与告警策略;",
    "ESKF 与 Nav2 定位的深度融合,以 ESKF 输出参与 AMCL 重定位评估,形成独立于 robot_localization 的定位链路。",
]
for i, it in enumerate(items, 1):
    para(f"{i}. {it}")

doc.save(OUT)
print("saved:", OUT)
